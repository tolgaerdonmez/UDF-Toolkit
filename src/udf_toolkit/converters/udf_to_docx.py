import base64
import io
import os
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_UNDERLINE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ._udf_io import extract_content_text, load_udf_xml


def get_alignment_style(alignment_value):
    if alignment_value == "1":
        return WD_ALIGN_PARAGRAPH.CENTER
    elif alignment_value == "3":
        return WD_ALIGN_PARAGRAPH.JUSTIFY
    elif alignment_value == "2":
        return WD_ALIGN_PARAGRAPH.RIGHT
    else:
        return WD_ALIGN_PARAGRAPH.LEFT


def convert_color(color_value):
    if color_value is None:
        return None

    try:
        color_int = int(color_value)
        if color_int < 0:
            color_int = 0xFFFFFFFF + color_int + 1

        r = (color_int >> 16) & 0xFF
        g = (color_int >> 8) & 0xFF
        b = color_int & 0xFF

        return (RGBColor(r, g, b), (r, g, b))
    except (ValueError, TypeError):
        return None


def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    run._r.append(fldChar1)

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    run._r.append(instrText)

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar2)


def set_cell_background(cell, color_info):
    color_obj, rgb_values = color_info
    shading_elm = OxmlElement('w:shd')
    hex_color = f"{rgb_values[0]:02X}{rgb_values[1]:02X}{rgb_values[2]:02X}"
    shading_elm.set(qn('w:fill'), hex_color)
    cell._tc.get_or_add_tcPr().append(shading_elm)


def process_background_image(document, bg_image_data, bg_image_source, output_file):
    if bg_image_data:
        try:
            image_bytes = base64.b64decode(bg_image_data)

            temp_img_path = os.path.join(
                os.path.dirname(output_file),
                os.path.splitext(os.path.basename(output_file))[0] + "_background.png",
            )
            with open(temp_img_path, "wb") as img_file:
                img_file.write(image_bytes)

            print(
                f"Background image saved to {temp_img_path}. "
                "Please manually set it as document background in Word."
            )
            return True
        except Exception as exc:
            print(f"Error processing background image data: {exc}")
    elif bg_image_source:
        print(
            f"Background image source path: {bg_image_source}. "
            "Please manually set it as document background in Word."
        )
    return False


def convert(udf_file: str, docx_file: str) -> None:
    root = load_udf_xml(udf_file)

    document = Document()

    for section in document.sections:
        section.different_first_page = False
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

    content_text = extract_content_text(root)

    properties_element = root.find('properties')
    page_format = properties_element.find('pageFormat') if properties_element is not None else None

    if page_format is not None:
        left_margin = float(page_format.get('leftMargin', '42.5')) / 72 * Inches(1).pt
        right_margin = float(page_format.get('rightMargin', '42.5')) / 72 * Inches(1).pt
        top_margin = float(page_format.get('topMargin', '42.5')) / 72 * Inches(1).pt
        bottom_margin = float(page_format.get('bottomMargin', '42.5')) / 72 * Inches(1).pt

        for section in document.sections:
            section.left_margin = Pt(left_margin)
            section.right_margin = Pt(right_margin)
            section.top_margin = Pt(top_margin)
            section.bottom_margin = Pt(bottom_margin)

            orientation = page_format.get('paperOrientation', '1')
            if orientation == '2':
                section.orientation = WD_ORIENT.LANDSCAPE
            else:
                section.orientation = WD_ORIENT.PORTRAIT

    if properties_element is not None:
        bg_image_elem = properties_element.find('bgImage')
        if bg_image_elem is not None:
            bg_image_data = bg_image_elem.get('bgImageData')
            bg_image_source = bg_image_elem.get('bgImageSource')
            process_background_image(document, bg_image_data, bg_image_source, docx_file)

    elements_element = root.find('elements')
    if elements_element is not None:
        header_element = elements_element.find('header')
        footer_element = elements_element.find('footer')

        if header_element is not None:
            section = document.sections[0]
            header = section.header

            for p in header.paragraphs:
                p._element.getparent().remove(p._element)
                p._p = None
                p._element = None

            header_para = header.add_paragraph()
            for elem in header_element:
                if elem.tag == 'content':
                    start_offset = int(elem.get('startOffset', 0))
                    length = int(elem.get('length', 0))
                    text = content_text[start_offset:start_offset+length]
                    header_para.add_run(text)

        if footer_element is not None:
            section = document.sections[0]
            footer = section.footer

            for p in footer.paragraphs:
                p._element.getparent().remove(p._element)
                p._p = None
                p._element = None

            footer_para = footer.add_paragraph()
            for elem in footer_element:
                if elem.tag == 'content':
                    start_offset = int(elem.get('startOffset', 0))
                    length = int(elem.get('length', 0))
                    text = content_text[start_offset:start_offset+length]
                    footer_para.add_run(text)

                elif elem.tag == 'pageNumber':
                    add_page_number(footer_para)

        for elem in elements_element:
            if elem.tag == 'paragraph':
                paragraph = document.add_paragraph()

                alignment = elem.get('Alignment', '0')
                paragraph.alignment = get_alignment_style(alignment)

                for child in elem:
                    if child.tag == 'content':
                        start_offset = int(child.get('startOffset', 0))
                        length = int(child.get('length', 0))
                        text = content_text[start_offset:start_offset+length]

                        run = paragraph.add_run(text)

                        font_family = child.get('family')
                        if font_family:
                            run.font.name = font_family

                        font_size = child.get('size')
                        if font_size:
                            run.font.size = Pt(float(font_size))

                        if child.get('bold', 'false') == 'true':
                            run.bold = True
                        if child.get('italic', 'false') == 'true':
                            run.italic = True

                        underline = child.get('underline')
                        if underline and underline != 'none':
                            run.underline = WD_UNDERLINE.SINGLE

                        color_info = convert_color(child.get('foreground'))
                        if color_info:
                            run.font.color.rgb = color_info[0]

                    elif child.tag == 'space':
                        paragraph.add_run(' ')

                    elif child.tag == 'tab':
                        paragraph.add_run('\t')

                    elif child.tag == 'image':
                        image_data = child.get('imageData')
                        if image_data:
                            image_bytes = base64.b64decode(image_data)
                            image_stream = io.BytesIO(image_bytes)
                            paragraph.add_run().add_picture(image_stream)

            elif elem.tag == 'table':
                column_count = int(elem.get('columnCount', 1))
                table = document.add_table(rows=0, cols=column_count)

                for row_elem in elem.findall('row'):
                    row_cells = table.add_row().cells
                    cell_elems = row_elem.findall('cell')

                    for i, cell_elem in enumerate(cell_elems):
                        cell = row_cells[i]
                        for para_elem in cell_elem.findall('paragraph'):
                            para = cell.add_paragraph()
                            alignment = para_elem.get('Alignment', '0')
                            para.alignment = get_alignment_style(alignment)

                            for child in para_elem:
                                if child.tag == 'content':
                                    start_offset = int(child.get('startOffset', 0))
                                    length = int(child.get('length', 0))
                                    text = content_text[start_offset:start_offset+length]

                                    run = para.add_run(text)

                                    font_family = child.get('family')
                                    if font_family:
                                        run.font.name = font_family

                                    font_size = child.get('size')
                                    if font_size:
                                        run.font.size = Pt(float(font_size))

                                    if child.get('bold', 'false') == 'true':
                                        run.bold = True
                                    if child.get('italic', 'false') == 'true':
                                        run.italic = True

                                    color_info = convert_color(child.get('foreground'))
                                    if color_info:
                                        run.font.color.rgb = color_info[0]

                                elif child.tag == 'space':
                                    para.add_run(' ')

                        cell_bg = cell_elem.get('bgColor')
                        if cell_bg:
                            color_info = convert_color(cell_bg)
                            if color_info:
                                set_cell_background(cell, color_info)

    document.save(docx_file)
