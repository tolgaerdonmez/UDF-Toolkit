import zipfile
from pathlib import Path

import pytest
from docx import Document

from udf_toolkit import convert, detect_format, normalize_format


def create_minimal_udf(path: Path, text: str) -> None:
    udf_template = '''<?xml version="1.0" encoding="UTF-8" ?>
<template format_id="1.8">
<content><![CDATA[{content}]]></content>
<properties><pageFormat mediaSizeName="1" leftMargin="42.51968479156494" rightMargin="28.34645652770996" topMargin="14.17322826385498" bottomMargin="14.17322826385498" paperOrientation="1" headerFOffset="20.0" footerFOffset="20.0" /></properties>
<elements resolver="hvl-default">
{elements}
</elements>
<styles><style name="default" description="Geçerli" family="Dialog" size="12" bold="false" italic="false" foreground="-13421773" FONT_ATTRIBUTE_KEY="javax.swing.plaf.FontUIResource[family=Dialog,name=Dialog,style=plain,size=12]" /><style name="hvl-default" family="Times New Roman" size="12" description="Gövde" /></styles>
</template>'''

    elements = f'<paragraph Alignment="0" LeftIndent="0.0" RightIndent="0.0"><content startOffset="0" length="{len(text)}" /></paragraph>'
    udf_content = udf_template.format(content=text, elements=elements)

    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as zipf:
        zipf.writestr('content.xml', udf_content)


def test_detect_format():
    assert detect_format("file.docx") == "docx"
    assert detect_format("file.udf") == "udf"
    assert detect_format("file.md") == "md"
    assert normalize_format("markdown") == "md"


def test_docx_to_udf_roundtrip(tmp_path: Path):
    docx_path = tmp_path / "input.docx"
    doc = Document()
    doc.add_paragraph("Hello world")
    doc.save(docx_path)

    udf_path = convert(docx_path, target="udf")
    assert udf_path.exists()

    with zipfile.ZipFile(udf_path, 'r') as zipf:
        assert 'content.xml' in zipf.namelist()

    docx_out = convert(udf_path, target="docx")
    assert docx_out.exists()

    doc_out = Document(docx_out)
    text = "\n".join(p.text for p in doc_out.paragraphs)
    assert "Hello world" in text


def test_udf_to_md(tmp_path: Path):
    udf_path = tmp_path / "input.udf"
    create_minimal_udf(udf_path, "Hello")

    md_path = convert(udf_path, target="md")
    assert md_path.exists()
    assert "Hello" in md_path.read_text(encoding="utf-8")
